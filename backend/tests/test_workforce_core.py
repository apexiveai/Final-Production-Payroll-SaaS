from datetime import UTC, date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
import json
import logging
import sys
import csv
from io import StringIO
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.routes.auth import router as auth_router
from app.api.routes.companies import router as companies_router
from app.api.routes.employees import router as employees_router
import app.api.routes.employees as employees_module
from app.api.routes.operations import router as operations_router
from app.api.routes.payroll import router as payroll_router
from app.api.routes.regional_reports import router as regional_reports_router
from app.api.routes.sites import router as sites_router
from app.api.routes.users import router as users_router
import app.main as main_module
from app.db.session import Base, get_db
from app.models.company import Company
from app.models.site import Site
from app.models.user import User
from app.models.workforce import (
    Attendance,
    AuditLog,
    Employee,
    EmployeeDocument,
    PayrollItem,
    PayrollRun,
    Shift,
    ShiftAssignment,
)
from app.security.dependencies import DBSession, get_current_user
from app.security.password import hash_password, verify_password
from app.security.permissions import resolve_company_id
from app.api.routes.payroll import money, prorated_monthly_salary
import app.cli.purge_expired_employee_documents as purge_documents_cli


@pytest.fixture
def api_client():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    with session_factory() as db:
        first = Company(name="Company One", code="ONE")
        second = Company(name="Company Two", code="TWO")
        db.add_all((first, second))
        db.flush()
        company_admin = User(
            id=uuid4(),
            company_id=first.id,
            username="company_admin",
            email="admin-one@companyone.com",
            password_hash="not-used",
            full_name="Company Admin",
            role="COMPANY_ADMIN",
            is_active=True,
        )
        super_admin = User(
            id=uuid4(),
            company_id=None,
            username="super_admin",
            email="admin@example.com",
            password_hash="not-used",
            full_name="Super Admin",
            role="SUPER_ADMIN",
            is_active=True,
        )
        db.add_all((company_admin, super_admin))
        db.commit()

    app = FastAPI()
    app.include_router(auth_router, prefix="/api/v1")
    app.include_router(companies_router, prefix="/api/v1")
    app.include_router(sites_router, prefix="/api/v1")
    app.include_router(employees_router, prefix="/api/v1")
    app.include_router(operations_router, prefix="/api/v1")
    app.include_router(payroll_router, prefix="/api/v1")
    app.include_router(regional_reports_router, prefix="/api/v1")
    app.include_router(users_router, prefix="/api/v1")

    def db_override():
        with session_factory() as session:
            yield session

    app.dependency_overrides[get_db] = db_override
    app.dependency_overrides[get_current_user] = lambda: company_admin
    client = TestClient(app)
    try:
        yield client, first, second, company_admin, super_admin, app, session_factory
    finally:
        client.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_company_admin_listing_is_tenant_scoped(api_client):
    client, first, second, *_ = api_client
    response = client.get("/api/v1/companies")
    assert response.status_code == 200
    assert [item["id"] for item in response.json()] == [str(first.id)]
    assert response.json()[0]["id"] != str(second.id)


def test_company_admin_cannot_read_another_company(api_client):
    client, _first, second, *_ = api_client
    response = client.get(f"/api/v1/companies/{second.id}")
    assert response.status_code == 403


def test_site_administrative_fields_create_update_and_validate(api_client):
    client, company, *_ = api_client
    created = client.post(
        "/api/v1/sites",
        json={
            "name": "Yangon Office",
            "code": "YGN-01",
            "region": "  Yangon Region  ",
            "district": "Yangon",
            "township": "Dagon",
            "ward_village_tract": "Ward 1",
        },
    )
    assert created.status_code == 201
    assert created.json()["region"] == "Yangon Region"
    assert created.json()["district"] == "Yangon"
    assert created.json()["township"] == "Dagon"
    assert created.json()["ward_village_tract"] == "Ward 1"
    updated = client.patch(
        f"/api/v1/sites/{created.json()['id']}",
        json={"township": "  Bahan  ", "ward_village_tract": None},
    )
    assert updated.status_code == 200
    assert updated.json()["township"] == "Bahan"
    assert updated.json()["ward_village_tract"] is None
    invalid = client.post(
        "/api/v1/sites",
        json={"name": "Blank Area", "code": "BLANK", "region": "   "},
    )
    assert invalid.status_code == 422


def test_regional_report_aggregates_scoped_data_and_comparable_insights(api_client):
    client, company, other_company, admin, super_admin, app, session_factory = api_client
    with session_factory() as db:
        primary_site = Site(
            company_id=company.id,
            name="Yangon Central",
            code="YGN-C",
            region="Yangon Region",
            district="Yangon",
            township="Dagon",
        )
        other_site = Site(
            company_id=other_company.id,
            name="Mandalay Branch",
            code="MDY-B",
            region="Mandalay Region",
        )
        empty_site = Site(
            company_id=company.id,
            name="No Records Office",
            code="NO-RECORDS",
            region="Sagaing Region",
        )
        db.add_all((primary_site, other_site, empty_site))
        db.flush()
        employee = Employee(
            company_id=company.id,
            site_id=primary_site.id,
            employee_code="REG-001",
            full_name="Regional Employee",
            hire_date=date(2025, 1, 1),
        )
        other_employee = Employee(
            company_id=other_company.id,
            site_id=other_site.id,
            employee_code="REG-002",
            full_name="Other Company Employee",
            hire_date=date(2025, 1, 1),
        )
        db.add_all((employee, other_employee))
        db.flush()
        for work_date, attendance_status, overtime in (
            (date(2026, 4, 1), "PRESENT", 0),
            (date(2026, 4, 2), "PRESENT", 0),
            (date(2026, 4, 3), "PRESENT", 30),
            (date(2026, 4, 4), "ABSENT", 0),
        ):
            db.add(
                Attendance(
                    employee_id=employee.id,
                    work_date=work_date,
                    status=attendance_status,
                    overtime_minutes=overtime,
                )
            )

        def add_run(first, last, status, gross, overtime):
            run = PayrollRun(
                company_id=company.id,
                period_start=first,
                period_end=last,
                currency="MMK",
                status=status,
                created_by=admin.id,
            )
            db.add(run)
            db.flush()
            db.add(
                PayrollItem(
                    run_id=run.id,
                    employee_id=employee.id,
                    base_pay=Decimal("100"),
                    overtime_pay=Decimal(overtime),
                    allowances=Decimal("0"),
                    bonus_pay=Decimal("0"),
                    deductions=Decimal("0"),
                    unpaid_leave_deduction=Decimal("0"),
                    gross_pay=Decimal(gross),
                    net_pay=Decimal(gross),
                )
            )

        add_run(date(2026, 4, 1), date(2026, 4, 2), "LOCKED", "100", "6")
        add_run(date(2026, 4, 3), date(2026, 4, 4), "PAID", "200", "10")
        add_run(date(2026, 4, 3), date(2026, 4, 4), "DRAFT", "900", "90")
        db.commit()
        primary_site_id = primary_site.id
        other_site_id = other_site.id
        empty_site_id = empty_site.id

    response = client.get(
        "/api/v1/reports/regional-workforce",
        params={
            "starts_on": "2026-04-03",
            "ends_on": "2026-04-04",
            "grouping_level": "region",
        },
    )
    assert response.status_code == 200
    report = response.json()
    assert [item["name"] for item in report["regions"]] == [
        "Sagaing Region",
        "Yangon Region",
    ]
    metric = next(item for item in report["regions"] if item["name"] == "Yangon Region")
    assert metric["site_ids"] == [str(primary_site_id)]
    assert metric["active_employee_count"] == 1
    assert metric["data_available"]["active_employee_count"] is True
    assert metric["data_available"]["attendance_rate"] is True
    assert metric["data_available"]["labor_cost"] is True
    assert metric["attendance_record_count"] == 2
    assert metric["attendance_rate"] == 50.0
    assert metric["overtime_minutes"] == 30
    assert metric["labor_cost"] == 200
    assert metric["overtime_cost"] == 10
    assert metric["currency"] == "MMK"
    unavailable = next(
        item for item in report["regions"] if item["name"] == "Sagaing Region"
    )
    assert unavailable["active_employee_count"] == 0
    assert unavailable["data_available"]["active_employee_count"] is True
    assert unavailable["data_available"]["attendance"] is False
    assert unavailable["data_available"]["labor_cost"] is False
    assert unavailable["attendance_record_count"] is None
    assert unavailable["attendance_data_available"] is False
    assert unavailable["attendance_rate"] is None
    assert unavailable["overtime_minutes"] is None
    assert unavailable["payroll_data_available"] is False
    assert unavailable["labor_cost"] is None
    assert all(
        insight["source"] == "analytics-derived" and insight["is_forecast"] is False
        for insight in report["insights"]
    )
    assert {insight["metric"] for insight in report["insights"]} == {
        "attendance_rate",
        "overtime_cost",
    }

    denied = client.get(
        "/api/v1/reports/regional-workforce",
        params={
            "starts_on": "2026-04-03",
            "ends_on": "2026-04-04",
            "company_id": str(other_company.id),
        },
    )
    assert denied.status_code == 403

    app.dependency_overrides[get_current_user] = lambda: super_admin
    aggregate = client.get(
        "/api/v1/reports/regional-workforce",
        params={"starts_on": "2026-04-03", "ends_on": "2026-04-04"},
    )
    assert aggregate.status_code == 200
    assert {item["key"] for item in aggregate.json()["sites"]} == {
        str(primary_site_id),
        str(other_site_id),
        str(empty_site_id),
    }

    supervisor = User(
        company_id=company.id,
        site_id=primary_site_id,
        username="site_supervisor",
        email="supervisor@example.com",
        password_hash="not-used",
        full_name="Site Supervisor",
        role="SUPERVISOR",
        is_active=True,
    )
    app.dependency_overrides[get_current_user] = lambda: supervisor
    supervisor_report = client.get(
        "/api/v1/reports/regional-workforce",
        params={"starts_on": "2026-04-03", "ends_on": "2026-04-04"},
    )
    assert supervisor_report.status_code == 200
    assert [item["key"] for item in supervisor_report.json()["sites"]] == [
        str(primary_site_id)
    ]
    denied_site = client.get(
        "/api/v1/reports/regional-workforce",
        params={
            "starts_on": "2026-04-03",
            "ends_on": "2026-04-04",
            "site_id": str(empty_site_id),
        },
    )
    assert denied_site.status_code == 403


def test_super_admin_company_list_supports_search_and_pagination(api_client):
    client, first, second, _company_admin, super_admin, app, _session_factory = api_client
    app.dependency_overrides[get_current_user] = lambda: super_admin
    response = client.get("/api/v1/companies", params={"search": "Two", "limit": 1})
    assert response.status_code == 200
    assert [item["id"] for item in response.json()] == [str(second.id)]
    response = client.get("/api/v1/companies", params={"offset": 1, "limit": 1})
    assert response.status_code == 200
    assert len(response.json()) == 1
    assert response.json()[0]["id"] in {str(first.id), str(second.id)}


def test_company_scope_rejects_cross_tenant_id(api_client):
    _client, first, second, company_admin, *_ = api_client
    assert resolve_company_id(company_admin, first.id) == first.id
    with pytest.raises(HTTPException) as exc:
        resolve_company_id(company_admin, second.id)
    assert exc.value.status_code == 403


def test_company_admin_can_create_company_user_but_not_company_admin(api_client):
    client, company, *_ = api_client
    response = client.post(
        "/api/v1/users",
        json={
            "company_id": str(company.id),
            "username": "hr_manager",
            "email": "hr@companyone.com",
            "full_name": "HR Manager",
            "password": "a-strong-password",
            "role": "HR",
        },
    )
    assert response.status_code == 201, response.text
    assert response.json()["company_id"] == str(company.id)
    assert response.json()["role"] == "HR"
    assert "password_hash" not in response.json()

    response = client.post(
        "/api/v1/users",
        json={
            "company_id": str(company.id),
            "username": "another_admin",
            "email": "another-admin@companyone.com",
            "full_name": "Another Admin",
            "password": "a-strong-password",
            "role": "COMPANY_ADMIN",
        },
    )
    assert response.status_code == 403


def test_company_admin_can_list_and_deactivate_only_company_users(api_client):
    client, company, second_company, _admin, _super_admin, _app, session_factory = api_client
    create_response = client.post(
        "/api/v1/users",
        json={
            "company_id": str(company.id),
            "username": "payroll_user",
            "email": "payroll@companyone.com",
            "full_name": "Payroll User",
            "password": "a-strong-password",
            "role": "PAYROLL",
        },
    )
    assert create_response.status_code == 201, create_response.text
    user_id = create_response.json()["id"]

    list_response = client.get("/api/v1/users")
    assert list_response.status_code == 200
    assert all(user["company_id"] == str(company.id) for user in list_response.json())

    cross_tenant_response = client.get(
        "/api/v1/users",
        params={"company_id": str(second_company.id)},
    )
    assert cross_tenant_response.status_code == 403

    status_response = client.patch(
        f"/api/v1/users/{user_id}/status",
        json={"is_active": False},
    )
    assert status_response.status_code == 200
    assert status_response.json()["is_active"] is False
    with session_factory() as db:
        managed_user = db.query(User).filter(User.id == UUID(user_id)).one()
        assert managed_user.token_version == 1


def test_supervisor_user_is_limited_to_assigned_site(api_client):
    client, company, _second, _company_admin, super_admin, app, session_factory = api_client
    with session_factory() as db:
        first_site = Site(company_id=company.id, name="North", code="NORTH")
        second_site = Site(company_id=company.id, name="South", code="SOUTH")
        db.add_all((first_site, second_site))
        db.flush()
        north_employee = Employee(
            company_id=company.id,
            site_id=first_site.id,
            employee_code="N-001",
            full_name="North Employee",
            hire_date=date(2026, 1, 1),
        )
        south_employee = Employee(
            company_id=company.id,
            site_id=second_site.id,
            employee_code="S-001",
            full_name="South Employee",
            hire_date=date(2026, 1, 1),
        )
        db.add_all((north_employee, south_employee))
        shift = Shift(
            company_id=company.id,
            name="Day",
            shift_type="DAY",
            starts_at=time(8),
            ends_at=time(17),
        )
        db.add(shift)
        db.flush()
        db.add_all(
            (
                ShiftAssignment(
                    employee_id=north_employee.id,
                    shift_id=shift.id,
                    work_date=date(2026, 1, 2),
                ),
                ShiftAssignment(
                    employee_id=south_employee.id,
                    shift_id=shift.id,
                    work_date=date(2026, 1, 2),
                ),
            )
        )
        db.commit()
        south_employee_id = south_employee.id

    app.dependency_overrides[get_current_user] = lambda: super_admin
    create_response = client.post(
        "/api/v1/users",
        json={
            "company_id": str(company.id),
            "site_id": str(first_site.id),
            "username": "north_supervisor",
            "email": "north-supervisor@companyone.com",
            "full_name": "North Supervisor",
            "password": "a-strong-password",
            "role": "SUPERVISOR",
        },
    )
    assert create_response.status_code == 201, create_response.text
    supervisor = User(
        id=UUID(create_response.json()["id"]),
        company_id=company.id,
        site_id=first_site.id,
        username="north_supervisor",
        email="north-supervisor@companyone.com",
        password_hash="not-used",
        full_name="North Supervisor",
        role="SUPERVISOR",
        is_active=True,
    )
    app.dependency_overrides[get_current_user] = lambda: supervisor

    response = client.get("/api/v1/employees")
    assert response.status_code == 200
    assert [employee["employee_code"] for employee in response.json()] == ["N-001"]
    response = client.get(
        "/api/v1/employees",
        params={"site_id": str(second_site.id)},
    )
    assert response.status_code == 403
    response = client.post(
        "/api/v1/attendance",
        json={
            "employee_id": str(south_employee_id),
            "work_date": "2026-01-01",
            "status": "PRESENT",
        },
    )
    assert response.status_code == 403
    response = client.get("/api/v1/shifts/assignments")
    assert response.status_code == 200
    assert [assignment["employee_code"] for assignment in response.json()] == ["N-001"]
    response = client.get(
        "/api/v1/reports/attendance.csv",
        params={
            "starts_on": "2026-01-01",
            "ends_on": "2026-01-31",
            "site_id": str(second_site.id),
        },
    )
    assert response.status_code == 403


def test_attendance_csv_export_escapes_spreadsheet_formulas(api_client):
    client, company, *_ = api_client
    with api_client[-1]() as db:
        employee = Employee(
            company_id=company.id,
            employee_code="=CMD",
            full_name="+Formula Employee",
            hire_date=date(2026, 1, 1),
        )
        db.add(employee)
        db.flush()
        db.add(
            Attendance(
                employee_id=employee.id,
                work_date=date(2026, 1, 2),
                status="PRESENT",
            )
        )
        db.commit()

    response = client.get(
        "/api/v1/reports/attendance.csv",
        params={"starts_on": "2026-01-01", "ends_on": "2026-01-31"},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "'=CMD" in response.text
    assert "'+Formula Employee" in response.text


def test_password_change_updates_hash_and_invalidates_existing_tokens(api_client):
    client, _company, _second, company_admin, _super_admin, app, session_factory = api_client
    old_password = "old-password-123"
    new_password = "new-password-456"
    with session_factory() as db:
        persisted_user = db.query(User).filter(User.id == company_admin.id).one()
        persisted_user.password_hash = hash_password(old_password)
        db.commit()

    def attached_user(db: DBSession):
        return db.query(User).filter(User.id == company_admin.id).one()

    app.dependency_overrides[get_current_user] = attached_user
    response = client.post(
        "/api/v1/auth/change-password",
        json={
            "current_password": old_password,
            "new_password": new_password,
        },
    )
    assert response.status_code == 200
    assert response.json()["reauthentication_required"] is True
    with session_factory() as db:
        persisted_user = db.query(User).filter(User.id == company_admin.id).one()
        assert persisted_user.token_version == 1
        assert verify_password(new_password, persisted_user.password_hash)


@pytest.mark.parametrize(
    ("first", "last", "expected"),
    [
        (date(2026, 1, 1), date(2026, 1, 31), Decimal("310000.00")),
        (date(2026, 1, 1), date(2026, 1, 15), Decimal("150000.00")),
        (date(2026, 1, 31), date(2026, 2, 1), Decimal("21071.43")),
    ],
)
def test_monthly_salary_proration(first, last, expected):
    assert prorated_monthly_salary(Decimal("310000"), first, last) == expected


def test_money_rounds_half_up():
    assert money(Decimal("1.005")) == Decimal("1.01")


def test_payroll_draft_calculates_monthly_salary_and_overtime(api_client):
    client, company, _second, _admin, _super_admin, _app, session_factory = api_client
    employee_id = uuid4()
    with session_factory() as db:
        employee = Employee(
            id=employee_id,
            company_id=company.id,
            employee_code="EMP-001",
            full_name="Test Employee",
            hire_date=date(2026, 1, 1),
            employment_status="ACTIVE",
            contract_type="PERMANENT",
            base_salary=Decimal("310000"),
            monthly_allowances=Decimal("31000"),
            monthly_bonus=Decimal("0"),
            monthly_deductions=Decimal("0"),
            currency="MMK",
        )
        db.add(employee)
        db.flush()
        db.add(
            Attendance(
                employee_id=employee.id,
                work_date=date(2026, 1, 15),
                status="PRESENT",
                overtime_minutes=60,
            )
        )
        db.commit()

    response = client.post(
        "/api/v1/payroll/runs",
        json={
            "period_start": "2026-01-01",
            "period_end": "2026-01-31",
            "currency": "MMK",
        },
    )
    assert response.status_code == 201
    run_id = response.json()["id"]
    assert response.json()["status"] == "DRAFT"
    items_response = client.get(f"/api/v1/payroll/runs/{run_id}/items")
    assert items_response.status_code == 200
    item = items_response.json()[0]
    assert Decimal(item["base_pay"]) == Decimal("310000.00")
    assert Decimal(item["overtime_pay"]) == Decimal("1875.00")
    assert Decimal(item["allowances"]) == Decimal("31000.00")
    assert Decimal(item["gross_pay"]) == Decimal("342875.00")
    assert Decimal(item["net_pay"]) == Decimal("342875.00")


def test_company_can_run_separate_payrolls_for_multiple_currencies(api_client):
    client, company, _second, _admin, _super_admin, _app, session_factory = api_client
    with session_factory() as db:
        mmk_employee = Employee(
            company_id=company.id,
            employee_code="MMK-001",
            full_name="MMK Employee",
            hire_date=date(2026, 1, 1),
            currency="MMK",
            base_salary=Decimal("310000"),
        )
        usd_employee = Employee(
            company_id=company.id,
            employee_code="USD-001",
            full_name="USD Employee",
            hire_date=date(2026, 1, 1),
            currency="USD",
            base_salary=Decimal("1000"),
        )
        db.add_all(
            [mmk_employee, usd_employee]
        )
        db.commit()
        mmk_employee_id = str(mmk_employee.id)
        usd_employee_id = str(usd_employee.id)

    payload = {
        "period_start": "2026-01-01",
        "period_end": "2026-01-31",
    }
    mmk_run = client.post("/api/v1/payroll/runs", json={**payload, "currency": "MMK"})
    usd_run = client.post("/api/v1/payroll/runs", json={**payload, "currency": "USD"})
    duplicate_mmk = client.post(
        "/api/v1/payroll/runs", json={**payload, "currency": "MMK"}
    )

    assert mmk_run.status_code == 201
    assert usd_run.status_code == 201
    assert duplicate_mmk.status_code == 409
    assert mmk_run.json()["currency"] == "MMK"
    assert usd_run.json()["currency"] == "USD"
    mmk_items = client.get(
        f"/api/v1/payroll/runs/{mmk_run.json()['id']}/items"
    ).json()
    usd_items = client.get(
        f"/api/v1/payroll/runs/{usd_run.json()['id']}/items"
    ).json()
    assert [item["employee_id"] for item in mmk_items] == [mmk_employee_id]
    assert [item["employee_id"] for item in usd_items] == [usd_employee_id]


def test_employee_csv_import_is_atomic_and_export_is_formula_safe(api_client):
    client, company, _second, _admin, _super_admin, _app, session_factory = api_client
    csv_body = (
        "employee_code,full_name,hire_date,base_salary,currency,email\n"
        "CSV-001,=FORMULA,2026-01-01,250000,MMK,employee@example.com\n"
        "CSV-002,Valid Employee,2026-02-01,200000,MMK,\n"
    )
    response = client.post(
        f"/api/v1/employees/import.csv?company_id={company.id}",
        files={"file": ("employees.csv", csv_body.encode(), "text/csv")},
    )
    assert response.status_code == 201
    assert response.json() == {"imported": 2}

    export = client.get(f"/api/v1/employees/export.csv?company_id={company.id}")
    assert export.status_code == 200
    exported_rows = list(csv.DictReader(StringIO(export.text)))
    assert exported_rows[0]["full_name"] == "'=FORMULA"
    assert exported_rows[0]["employee_code"] == "CSV-001"

    created = client.post(
        "/api/v1/employees",
        json={
            "employee_code": "DIRECT-001",
            "full_name": "Direct Create",
            "hire_date": "2026-04-01",
            "base_salary": "150000.00",
            "currency": "MMK",
        },
    )
    assert created.status_code == 201
    assert created.json()["employee_code"] == "DIRECT-001"
    assert created.json()["full_name"] == "Direct Create"

    duplicate_and_new_rows = (
        "employee_code,full_name,hire_date,base_salary,currency\n"
        "CSV-003,Would Be New,2026-03-01,100000,MMK\n"
        "CSV-001,Duplicate,2026-03-01,100000,MMK\n"
    )
    rejected = client.post(
        f"/api/v1/employees/import.csv?company_id={company.id}",
        files={"file": ("employees.csv", duplicate_and_new_rows.encode(), "text/csv")},
    )
    assert rejected.status_code == 409
    with session_factory() as db:
        codes = {
            code
            for (code,) in db.query(Employee.employee_code).filter(
                Employee.company_id == company.id
            ).all()
        }
        assert codes == {"CSV-001", "CSV-002", "DIRECT-001"}
        assert db.query(AuditLog).filter(
            AuditLog.action == "employee.csv_imported"
        ).count() == 1


def test_locked_payroll_excel_export_contains_safe_text_and_pay_values(api_client):
    client, company, _second, company_admin, _super_admin, _app, session_factory = api_client
    with session_factory() as db:
        employee = Employee(
            company_id=company.id,
            employee_code="=FORMULA",
            full_name="+Spreadsheet Name",
            hire_date=date(2026, 1, 1),
        )
        db.add(employee)
        db.flush()
        run = PayrollRun(
            company_id=company.id,
            period_start=date(2026, 1, 1),
            period_end=date(2026, 1, 31),
            currency="MMK",
            status="LOCKED",
            created_by=company_admin.id,
        )
        db.add(run)
        db.flush()
        db.add(
            PayrollItem(
                run_id=run.id,
                employee_id=employee.id,
                base_pay=Decimal("100000.00"),
                overtime_pay=Decimal("5000.00"),
                allowances=Decimal("10000.00"),
                bonus_pay=Decimal("0.00"),
                deductions=Decimal("1000.00"),
                unpaid_leave_deduction=Decimal("0.00"),
                gross_pay=Decimal("115000.00"),
                net_pay=Decimal("114000.00"),
            )
        )
        db.commit()
        run_id = run.id

    csv_response = client.get(f"/api/v1/payroll/runs/{run_id}/export.csv")
    assert csv_response.status_code == 200
    csv_rows = list(csv.DictReader(StringIO(csv_response.text)))
    assert csv_rows[0]["employee_code"] == "'=FORMULA"
    assert csv_rows[0]["employee_name"] == "'+Spreadsheet Name"

    accounting_response = client.get(
        f"/api/v1/payroll/runs/{run_id}/export.accounting.csv"
    )
    assert accounting_response.status_code == 200
    journal_rows = list(csv.DictReader(StringIO(accounting_response.text)))
    assert [row["account_label"] for row in journal_rows] == [
        "Payroll expense",
        "Employee payable",
        "Deductions payable",
    ]
    debit_total = sum(Decimal(row["debit"]) for row in journal_rows)
    credit_total = sum(Decimal(row["credit"]) for row in journal_rows)
    assert debit_total == credit_total == Decimal("115000.00")

    response = client.get(f"/api/v1/payroll/runs/{run_id}/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    worksheet = workbook["Payroll"]
    assert worksheet["A2"].value == "=FORMULA"
    assert worksheet["A2"].data_type == "s"
    assert worksheet["B2"].value == "+Spreadsheet Name"
    assert worksheet["B2"].data_type == "s"
    assert worksheet["K2"].value == 114000
    with session_factory() as db:
        assert db.query(AuditLog).filter(
            AuditLog.action == "payroll.xlsx_exported"
        ).count() == 1


def test_readiness_endpoint_checks_database(api_client, monkeypatch):
    _client, *_values, app, session_factory = api_client
    monkeypatch.setattr(main_module, "SessionLocal", session_factory)
    app.add_api_route("/api/v1/health/ready", main_module.readiness, methods=["GET"])
    client = TestClient(app)
    try:
        response = client.get("/api/v1/health/ready")
    finally:
        client.close()
    assert response.status_code == 200
    assert response.json()["database"] == "connected"


def test_employee_document_upload_list_and_signed_download(api_client, monkeypatch):
    client, company, _second, _admin, _super_admin, _app, session_factory = api_client
    with session_factory() as db:
        employee = Employee(
            company_id=company.id,
            employee_code="DOC-001",
            full_name="Document Employee",
            hire_date=date(2026, 1, 1),
        )
        db.add(employee)
        db.commit()
        employee_id = employee.id

    class FakeObjectStore:
        def __init__(self):
            self.upload = None

        def put_object(self, **kwargs):
            self.upload = kwargs

        def generate_presigned_url(self, _operation, *, Params, ExpiresIn):
            assert ExpiresIn == 60
            assert Params["Key"] == self.upload["Key"]
            return "https://storage.example/download?temporary=signature"

    storage = FakeObjectStore()
    monkeypatch.setattr(employees_module, "get_object_storage_bucket", lambda: "private-docs")
    monkeypatch.setattr(employees_module, "get_object_store", lambda: storage)
    upload = client.post(
        f"/api/v1/employees/{employee_id}/documents",
        files={"file": ("identity.pdf", b"%PDF-1.7 sample", "application/pdf")},
        data={"category": "DOCUMENT"},
    )
    assert upload.status_code == 201
    assert "object_key" not in upload.json()
    assert storage.upload["Bucket"] == "private-docs"
    assert storage.upload["ServerSideEncryption"] == "AES256"

    listed = client.get(f"/api/v1/employees/{employee_id}/documents")
    assert listed.status_code == 200
    assert len(listed.json()) == 1
    assert listed.json()[0]["original_filename"] == "identity.pdf"

    downloaded = client.get(
        f"/api/v1/employees/{employee_id}/documents/{upload.json()['id']}/download"
    )
    assert downloaded.status_code == 200
    assert downloaded.json() == {
        "url": "https://storage.example/download?temporary=signature",
        "expires_in": 60,
    }
    with session_factory() as db:
        assert db.query(EmployeeDocument).count() == 1
        assert db.query(AuditLog).filter(
            AuditLog.action.in_(
                ["employee.document_uploaded", "employee.document_downloaded"]
            )
        ).count() == 2


def test_employee_document_upload_rejects_invalid_content_and_other_tenant(
    api_client, monkeypatch
):
    client, company, second_company, _admin, _super_admin, _app, session_factory = api_client
    with session_factory() as db:
        employee = Employee(
            company_id=company.id,
            employee_code="DOC-002",
            full_name="Document Employee",
            hire_date=date(2026, 1, 1),
        )
        other_employee = Employee(
            company_id=second_company.id,
            employee_code="DOC-003",
            full_name="Other Tenant",
            hire_date=date(2026, 1, 1),
        )
        db.add_all((employee, other_employee))
        db.commit()
        employee_id = employee.id
        other_employee_id = other_employee.id

    storage_calls = []
    monkeypatch.setattr(employees_module, "get_object_storage_bucket", lambda: "private-docs")
    monkeypatch.setattr(
        employees_module,
        "get_object_store",
        lambda: storage_calls.append("storage"),
    )
    invalid_upload = client.post(
        f"/api/v1/employees/{employee_id}/documents",
        files={"file": ("identity.pdf", b"not a PDF", "application/pdf")},
    )
    other_tenant_list = client.get(
        f"/api/v1/employees/{other_employee_id}/documents"
    )
    assert invalid_upload.status_code == 415
    assert other_tenant_list.status_code == 403
    assert storage_calls == []


def test_expired_document_cleanup_dry_run_and_audited_delete(api_client, monkeypatch, capsys):
    _client, company, _second, _admin, _super_admin, _app, session_factory = api_client
    with session_factory() as db:
        employee = Employee(
            company_id=company.id,
            employee_code="DOC-004",
            full_name="Expired Document Employee",
            hire_date=date(2026, 1, 1),
        )
        db.add(employee)
        db.flush()
        document = EmployeeDocument(
            employee_id=employee.id,
            category="DOCUMENT",
            original_filename="expired.pdf",
            content_type="application/pdf",
            object_key=f"companies/{company.id}/employees/{employee.id}/expired.pdf",
            size_bytes=128,
            created_at=datetime.now(UTC) - timedelta(days=2),
        )
        db.add(document)
        db.commit()
        document_id = document.id

    class FakeObjectStore:
        def __init__(self):
            self.deleted_keys = []

        def delete_object(self, *, Bucket, Key):
            assert Bucket == "private-docs"
            self.deleted_keys.append(Key)

    storage = FakeObjectStore()
    monkeypatch.setattr(purge_documents_cli, "SessionLocal", session_factory)
    monkeypatch.setattr(purge_documents_cli, "get_object_storage_bucket", lambda: "private-docs")
    monkeypatch.setattr(purge_documents_cli, "get_object_store", lambda: storage)
    monkeypatch.setattr(purge_documents_cli.settings, "employee_document_retention_days", 1)
    monkeypatch.setattr(sys, "argv", ["purge_expired_employee_documents", "--dry-run"])
    purge_documents_cli.main()
    assert "would be purged: 1" in capsys.readouterr().out
    assert storage.deleted_keys == []
    with session_factory() as db:
        assert db.get(EmployeeDocument, document_id) is not None

    monkeypatch.setattr(sys, "argv", ["purge_expired_employee_documents"])
    purge_documents_cli.main()
    assert "purged: 1" in capsys.readouterr().out
    assert len(storage.deleted_keys) == 1
    with session_factory() as db:
        assert db.get(EmployeeDocument, document_id) is None
        assert db.query(AuditLog).filter(
            AuditLog.action == "employee.documents_retention_purged"
        ).count() == 1


def test_api_request_logs_exclude_query_values_and_return_request_id(caplog):
    caplog.set_level(logging.INFO, logger="workforce.api")
    client = TestClient(main_module.app)
    try:
        response = client.get("/api/v1/health?marker=do-not-log-this")
    finally:
        client.close()

    assert response.status_code == 200
    assert response.headers["X-Request-ID"]
    records = [record for record in caplog.records if record.name == "workforce.api"]
    assert len(records) == 1
    assert "do-not-log-this" not in records[0].message
    log_payload = json.loads(records[0].message)
    assert log_payload["route"] == "/api/v1/health"
    assert log_payload["status_code"] == 200
