import calendar
import csv
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO, StringIO
from uuid import UUID

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func

from app.models.workforce import (
    Attendance,
    Employee,
    LeaveRequest,
    LeaveType,
    PayrollItem,
    PayrollRun,
)
from app.models.site import Site
from app.schemas.workforce import (
    PayrollItemResponse,
    PayrollItemAdjustment,
    PayrollRunCreate,
    PayrollRunResponse,
)
from app.security.dependencies import CurrentUser, DBSession
from app.security.permissions import PAYROLL_ROLES, require_role, resolve_company_id
from app.services.audit import record_audit
from app.services.csv_export import safe_csv_cell


router = APIRouter(tags=["Payroll and Reports"])
CENT = Decimal("0.01")


def money(value: Decimal) -> Decimal:
    return value.quantize(CENT, rounding=ROUND_HALF_UP)


def days_inclusive(first: date, last: date) -> int:
    return (last - first).days + 1


def prorated_monthly_salary(salary: Decimal, first: date, last: date) -> Decimal:
    total = Decimal("0")
    month_start = first.replace(day=1)
    while month_start <= last:
        month_end = month_start.replace(
            day=calendar.monthrange(month_start.year, month_start.month)[1]
        )
        overlap_start = max(first, month_start)
        overlap_end = min(last, month_end)
        if overlap_start <= overlap_end:
            total += salary * Decimal(days_inclusive(overlap_start, overlap_end)) / Decimal(
                calendar.monthrange(month_start.year, month_start.month)[1]
            )
        if month_start.month == 12:
            month_start = date(month_start.year + 1, 1, 1)
        else:
            month_start = date(month_start.year, month_start.month + 1, 1)
    return money(total)


def unpaid_leave_dates(db, employee_id: UUID, first: date, last: date) -> set[date]:
    requests = (
        db.query(LeaveRequest, LeaveType)
        .join(LeaveType, LeaveRequest.leave_type_id == LeaveType.id)
        .filter(
            LeaveRequest.employee_id == employee_id,
            LeaveRequest.status == "APPROVED",
            LeaveType.is_paid.is_(False),
            LeaveRequest.starts_on <= last,
            LeaveRequest.ends_on >= first,
        )
        .all()
    )
    days: set[date] = set()
    for leave, _ in requests:
        day = max(first, leave.starts_on)
        end = min(last, leave.ends_on)
        while day <= end:
            days.add(day)
            day = date.fromordinal(day.toordinal() + 1)
    return days


def get_run(db, run_id: UUID, current_user):
    run = db.query(PayrollRun).filter(PayrollRun.id == run_id).first()
    if not run:
        raise HTTPException(status_code=404, detail="Payroll run not found")
    resolve_company_id(current_user, run.company_id)
    return run


@router.post("/payroll/runs", response_model=PayrollRunResponse, status_code=201)
def create_payroll_run(
    payload: PayrollRunCreate,
    current_user: CurrentUser,
    db: DBSession,
):
    require_role(current_user, PAYROLL_ROLES)
    company_id = resolve_company_id(current_user, payload.company_id)
    duplicate = db.query(PayrollRun.id).filter(
        PayrollRun.company_id == company_id,
        PayrollRun.period_start == payload.period_start,
        PayrollRun.period_end == payload.period_end,
        PayrollRun.currency == payload.currency,
        PayrollRun.reverses_run_id.is_(None),
        PayrollRun.status != "REVERSED",
    ).first()
    if duplicate:
        raise HTTPException(status_code=409, detail="An active payroll run already exists for this period")
    employees = db.query(Employee).filter(
        Employee.company_id == company_id,
        Employee.employment_status == "ACTIVE",
        Employee.currency == payload.currency,
    ).order_by(Employee.id).all()
    if not employees:
        raise HTTPException(
            status_code=422,
            detail="No active employees use the payroll run currency",
        )

    run = PayrollRun(
        company_id=company_id,
        period_start=payload.period_start,
        period_end=payload.period_end,
        currency=payload.currency,
        status="DRAFT",
        created_by=current_user.id,
    )
    db.add(run)
    db.flush()

    period_days = days_inclusive(payload.period_start, payload.period_end)
    for employee in employees:
        base_pay = prorated_monthly_salary(
            employee.base_salary, payload.period_start, payload.period_end
        )
        allowance_pay = money(
            prorated_monthly_salary(
                employee.monthly_allowances,
                payload.period_start,
                payload.period_end,
            )
            + payload.allowances
        )
        bonus_pay = prorated_monthly_salary(
            employee.monthly_bonus, payload.period_start, payload.period_end
        )
        regular_deductions = money(
            prorated_monthly_salary(
                employee.monthly_deductions,
                payload.period_start,
                payload.period_end,
            )
            + payload.deductions
        )
        overtime_minutes = (
            db.query(func.coalesce(func.sum(Attendance.overtime_minutes), 0))
            .filter(
                Attendance.employee_id == employee.id,
                Attendance.work_date >= payload.period_start,
                Attendance.work_date <= payload.period_end,
                Attendance.status.notin_(("ABSENT", "LEAVE", "HOLIDAY")),
            )
            .scalar()
        )
        hours_in_period = Decimal(period_days) * payload.standard_hours_per_day
        hourly_rate = base_pay / hours_in_period if hours_in_period else Decimal("0")
        overtime_pay = money(
            Decimal(overtime_minutes or 0)
            / Decimal("60")
            * hourly_rate
            * payload.overtime_multiplier
        )
        unpaid_dates = unpaid_leave_dates(
            db, employee.id, payload.period_start, payload.period_end
        )
        absent_dates = {
            work_date
            for (work_date,) in db.query(Attendance.work_date).filter(
                Attendance.employee_id == employee.id,
                Attendance.work_date >= payload.period_start,
                Attendance.work_date <= payload.period_end,
                Attendance.status == "ABSENT",
            ).all()
        }
        unpaid_days = len(unpaid_dates | absent_dates)
        unpaid_deduction = money(
            base_pay * Decimal(unpaid_days) / Decimal(period_days)
        )
        gross_pay = money(base_pay + overtime_pay + allowance_pay + bonus_pay)
        net_pay = money(gross_pay - regular_deductions - unpaid_deduction)
        if net_pay < 0:
            raise HTTPException(
                status_code=422,
                detail=f"Payroll deductions exceed gross pay for employee {employee.employee_code}",
            )
        db.add(
            PayrollItem(
                run_id=run.id,
                employee_id=employee.id,
                base_pay=base_pay,
                overtime_pay=overtime_pay,
                allowances=allowance_pay,
                bonus_pay=bonus_pay,
                deductions=regular_deductions,
                unpaid_leave_deduction=unpaid_deduction,
                gross_pay=gross_pay,
                net_pay=net_pay,
            )
        )
    record_audit(
        db, current_user, "payroll.draft_created", "payroll_run", run.id,
        company_id=company_id,
    )
    db.commit()
    db.refresh(run)
    return run


@router.patch(
    "/payroll/runs/{run_id}/items/{employee_id}",
    response_model=PayrollItemResponse,
)
def adjust_payroll_item(
    run_id: UUID,
    employee_id: UUID,
    payload: PayrollItemAdjustment,
    current_user: CurrentUser,
    db: DBSession,
):
    require_role(current_user, PAYROLL_ROLES)
    run = get_run(db, run_id, current_user)
    if run.status != "DRAFT":
        raise HTTPException(status_code=409, detail="Only draft payroll items can be adjusted")
    item = db.query(PayrollItem).filter(
        PayrollItem.run_id == run.id,
        PayrollItem.employee_id == employee_id,
    ).first()
    if not item:
        raise HTTPException(status_code=404, detail="Payroll item not found")
    item.allowances = money(payload.allowances)
    item.bonus_pay = money(payload.bonus)
    item.deductions = money(payload.deductions)
    item.gross_pay = money(
        item.base_pay + item.overtime_pay + item.allowances + item.bonus_pay
    )
    item.net_pay = money(
        item.gross_pay - item.deductions - item.unpaid_leave_deduction
    )
    if item.net_pay < 0:
        raise HTTPException(status_code=422, detail="Deductions exceed gross pay")
    record_audit(db, current_user, "payroll.item_adjusted", "payroll_item", item.id, company_id=run.company_id)
    db.commit()
    db.refresh(item)
    return item


@router.get("/payroll/runs", response_model=list[PayrollRunResponse])
def list_payroll_runs(
    current_user: CurrentUser,
    db: DBSession,
    company_id: UUID | None = None,
    offset: int = Query(default=0, ge=0),
    limit: int = Query(default=50, ge=1, le=100),
):
    require_role(current_user, PAYROLL_ROLES)
    target_company = resolve_company_id(current_user, company_id)
    return db.query(PayrollRun).filter(
        PayrollRun.company_id == target_company
    ).order_by(PayrollRun.period_end.desc()).offset(offset).limit(limit).all()


@router.get("/payroll/payslips")
def list_my_payslips(
    current_user: CurrentUser,
    db: DBSession,
    offset: int = Query(default=0, ge=0),
    limit: int = Query(default=50, ge=1, le=100),
):
    if current_user.role != "EMPLOYEE":
        raise HTTPException(status_code=403, detail="Employee account required")
    employee = db.query(Employee).filter(Employee.user_id == current_user.id).first()
    if not employee:
        raise HTTPException(status_code=404, detail="Employee profile not found")
    rows = (
        db.query(PayrollItem, PayrollRun)
        .join(PayrollRun, PayrollItem.run_id == PayrollRun.id)
        .filter(
            PayrollItem.employee_id == employee.id,
            PayrollRun.status.in_(("LOCKED", "PAID")),
        )
        .order_by(PayrollRun.period_end.desc())
        .offset(offset)
        .limit(limit)
        .all()
    )
    return [
        {
            "run_id": str(run.id),
            "period_start": run.period_start,
            "period_end": run.period_end,
            "currency": run.currency,
            "status": run.status,
            "base_pay": item.base_pay,
            "overtime_pay": item.overtime_pay,
            "allowances": item.allowances,
            "bonus_pay": item.bonus_pay,
            "deductions": item.deductions,
            "unpaid_leave_deduction": item.unpaid_leave_deduction,
            "gross_pay": item.gross_pay,
            "net_pay": item.net_pay,
        }
        for item, run in rows
    ]


@router.get("/payroll/runs/{run_id}/items", response_model=list[PayrollItemResponse])
def list_payroll_items(
    run_id: UUID,
    current_user: CurrentUser,
    db: DBSession,
):
    if current_user.role == "EMPLOYEE":
        run = db.query(PayrollRun).filter(PayrollRun.id == run_id).first()
        if not run or run.status not in {"LOCKED", "PAID"}:
            raise HTTPException(status_code=404, detail="Payslip not found")
        employee = db.query(Employee).filter(Employee.user_id == current_user.id).first()
        if not employee:
            raise HTTPException(status_code=404, detail="Employee profile not found")
        return db.query(PayrollItem).filter(
            PayrollItem.run_id == run.id,
            PayrollItem.employee_id == employee.id,
        ).all()
    require_role(current_user, PAYROLL_ROLES)
    run = get_run(db, run_id, current_user)
    return db.query(PayrollItem).filter(PayrollItem.run_id == run.id).all()


@router.post("/payroll/runs/{run_id}/transition", response_model=PayrollRunResponse)
def transition_payroll(
    run_id: UUID,
    current_user: CurrentUser,
    db: DBSession,
    target_status: str = Query(pattern="^(REVIEWED|APPROVED|LOCKED)$"),
):
    run = get_run(db, run_id, current_user)
    allowed = {
        "REVIEWED": ({"DRAFT"}, PAYROLL_ROLES),
        "APPROVED": ({"REVIEWED"}, {"SUPER_ADMIN", "COMPANY_ADMIN"}),
        "LOCKED": ({"APPROVED"}, {"SUPER_ADMIN", "COMPANY_ADMIN"}),
    }
    previous_states, roles = allowed[target_status]
    require_role(current_user, roles)
    if run.status not in previous_states:
        raise HTTPException(
            status_code=409,
            detail=f"Cannot transition payroll from {run.status} to {target_status}",
        )
    if target_status == "APPROVED":
        if run.created_by == current_user.id or run.reviewed_by == current_user.id:
            raise HTTPException(status_code=403, detail="Payroll approval requires a different reviewer")
        run.approved_by = current_user.id
    if target_status == "REVIEWED":
        run.reviewed_by = current_user.id
    run.status = target_status
    record_audit(
        db, current_user, f"payroll.{target_status.lower()}", "payroll_run", run.id,
        company_id=run.company_id,
    )
    db.commit()
    db.refresh(run)
    return run


@router.post("/payroll/runs/{run_id}/payments")
def mark_payroll_paid(
    run_id: UUID,
    current_user: CurrentUser,
    db: DBSession,
):
    require_role(current_user, PAYROLL_ROLES)
    run = get_run(db, run_id, current_user)
    if run.status != "LOCKED":
        raise HTTPException(status_code=409, detail="Only locked payroll can be paid")
    db.query(PayrollItem).filter(
        PayrollItem.run_id == run.id,
        PayrollItem.payment_status == "UNPAID",
    ).update({PayrollItem.payment_status: "PAID"}, synchronize_session="fetch")
    run.status = "PAID"
    record_audit(
        db, current_user, "payroll.paid", "payroll_run", run.id,
        company_id=run.company_id,
    )
    db.commit()
    return {"run_id": str(run.id), "status": run.status}


@router.post("/payroll/runs/{run_id}/reverse", response_model=PayrollRunResponse)
def reverse_payroll(
    run_id: UUID,
    current_user: CurrentUser,
    db: DBSession,
    reason: str = Query(min_length=5, max_length=1000),
):
    require_role(current_user, {"SUPER_ADMIN", "COMPANY_ADMIN"})
    original = get_run(db, run_id, current_user)
    if original.status not in {"APPROVED", "LOCKED", "PAID"}:
        raise HTTPException(status_code=409, detail="Only approved or finalized payroll can be reversed")
    existing = db.query(PayrollRun.id).filter(
        PayrollRun.reverses_run_id == original.id
    ).first()
    if existing:
        raise HTTPException(status_code=409, detail="Payroll has already been reversed")
    reversal = PayrollRun(
        company_id=original.company_id,
        period_start=original.period_start,
        period_end=original.period_end,
        currency=original.currency,
        status="DRAFT",
        created_by=current_user.id,
        reverses_run_id=original.id,
        correction_reason=reason,
    )
    db.add(reversal)
    db.flush()
    for item in db.query(PayrollItem).filter(PayrollItem.run_id == original.id).all():
        db.add(
            PayrollItem(
                run_id=reversal.id,
                employee_id=item.employee_id,
                base_pay=-item.base_pay,
                overtime_pay=-item.overtime_pay,
                allowances=-item.allowances,
                bonus_pay=-item.bonus_pay,
                deductions=-item.deductions,
                unpaid_leave_deduction=-item.unpaid_leave_deduction,
                gross_pay=-item.gross_pay,
                net_pay=-item.net_pay,
            )
        )
    original.status = "REVERSED"
    record_audit(
        db, current_user, "payroll.reversed", "payroll_run", original.id,
        company_id=original.company_id,
    )
    record_audit(
        db, current_user, "payroll.reversal_draft_created", "payroll_run", reversal.id,
        company_id=original.company_id,
    )
    db.commit()
    db.refresh(reversal)
    return reversal


@router.get("/reports/summary")
def report_summary(
    current_user: CurrentUser,
    db: DBSession,
    company_id: UUID | None = None,
):
    if current_user.role == "EMPLOYEE":
        raise HTTPException(status_code=403, detail="Reports access denied")
    target_company = resolve_company_id(current_user, company_id)
    today = date.today()
    active_employees_query = db.query(func.count(Employee.id)).filter(
        Employee.company_id == target_company,
        Employee.employment_status == "ACTIVE",
    )
    active_sites_query = db.query(func.count(Site.id)).filter(
        Site.company_id == target_company,
        Site.is_active.is_(True),
    )
    attendance_query = (
        db.query(func.count(Attendance.id))
        .join(Employee)
        .filter(Employee.company_id == target_company, Attendance.work_date == today)
    )
    pending_leave = (
        db.query(func.count(LeaveRequest.id))
        .join(Employee)
        .filter(Employee.company_id == target_company, LeaveRequest.status == "PENDING")
    )
    if current_user.role == "SUPERVISOR":
        if current_user.site_id is None:
            raise HTTPException(status_code=403, detail="Supervisor site is not configured")
        active_employees_query = active_employees_query.filter(
            Employee.site_id == current_user.site_id
        )
        attendance_query = attendance_query.filter(
            Employee.site_id == current_user.site_id
        )
        pending_leave = pending_leave.filter(Employee.site_id == current_user.site_id)
        active_sites_query = active_sites_query.filter(Site.id == current_user.site_id)
    active_employees = active_employees_query.scalar()
    active_sites = active_sites_query.scalar()
    attendance_count = attendance_query.scalar()
    pending_leave = pending_leave.scalar()
    pending_payroll = db.query(func.count(PayrollRun.id)).filter(
        PayrollRun.company_id == target_company,
        PayrollRun.status.in_(("DRAFT", "REVIEWED")),
    ).scalar()
    if current_user.role == "SUPERVISOR":
        pending_payroll = 0
    return {
        "company_id": str(target_company),
        "active_employees": active_employees or 0,
        "active_sites": active_sites or 0,
        "attendance_records_today": attendance_count or 0,
        "pending_leave_approvals": pending_leave or 0,
        "pending_payroll_runs": pending_payroll or 0,
    }


@router.get("/payroll/runs/{run_id}/export.csv")
def export_payroll_csv(
    run_id: UUID,
    current_user: CurrentUser,
    db: DBSession,
):
    require_role(current_user, PAYROLL_ROLES)
    run = get_run(db, run_id, current_user)
    if run.status not in {"LOCKED", "PAID"}:
        raise HTTPException(status_code=409, detail="Only locked payroll can be exported")
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(
        ["employee_id", "employee_code", "employee_name", "currency", "gross_pay", "deductions", "net_pay", "payment_status"]
    )
    rows = (
        db.query(PayrollItem, Employee)
        .join(Employee, PayrollItem.employee_id == Employee.id)
        .filter(PayrollItem.run_id == run.id)
        .order_by(Employee.employee_code)
        .all()
    )
    for item, employee in rows:
        writer.writerow(
            [
                employee.id,
                safe_csv_cell(employee.employee_code),
                safe_csv_cell(employee.full_name),
                safe_csv_cell(run.currency),
                item.gross_pay,
                item.deductions + item.unpaid_leave_deduction,
                item.net_pay,
                safe_csv_cell(item.payment_status),
            ]
        )
    record_audit(
        db, current_user, "payroll.exported", "payroll_run", run.id,
        company_id=run.company_id,
    )
    db.commit()
    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="payroll-{run.id}.csv"'},
    )


@router.get("/payroll/runs/{run_id}/export.accounting.csv")
def export_payroll_accounting_csv(
    run_id: UUID,
    current_user: CurrentUser,
    db: DBSession,
):
    require_role(current_user, PAYROLL_ROLES)
    run = get_run(db, run_id, current_user)
    if run.status not in {"LOCKED", "PAID"}:
        raise HTTPException(status_code=409, detail="Only locked payroll can be exported")

    items = db.query(PayrollItem).filter(PayrollItem.run_id == run.id).all()
    gross_pay = sum((item.gross_pay for item in items), Decimal("0.00"))
    net_pay = sum((item.net_pay for item in items), Decimal("0.00"))
    deductions = sum(
        (item.deductions + item.unpaid_leave_deduction for item in items),
        Decimal("0.00"),
    )
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "payroll_run_id",
            "period_start",
            "period_end",
            "currency",
            "account_label",
            "debit",
            "credit",
        ]
    )
    journal_lines = (
        ("Payroll expense", gross_pay, True),
        ("Employee payable", net_pay, False),
        ("Deductions payable", deductions, False),
    )
    for account, amount, is_debit in journal_lines:
        debit = amount if is_debit else Decimal("0.00")
        credit = Decimal("0.00") if is_debit else amount
        if amount < 0:
            debit, credit = credit, debit
        writer.writerow(
            [
                run.id,
                run.period_start.isoformat(),
                run.period_end.isoformat(),
                safe_csv_cell(run.currency),
                safe_csv_cell(account),
                f"{debit:.2f}",
                f"{credit:.2f}",
            ]
        )

    record_audit(
        db,
        current_user,
        "payroll.accounting_exported",
        "payroll_run",
        run.id,
        company_id=run.company_id,
    )
    db.commit()
    return Response(
        content=output.getvalue(),
        media_type="text/csv",
        headers={
            "Content-Disposition": (
                f'attachment; filename="payroll-accounting-{run.id}.csv"'
            )
        },
    )


@router.get("/payroll/runs/{run_id}/export.xlsx")
def export_payroll_xlsx(
    run_id: UUID,
    current_user: CurrentUser,
    db: DBSession,
):
    require_role(current_user, PAYROLL_ROLES)
    run = get_run(db, run_id, current_user)
    if run.status not in {"LOCKED", "PAID"}:
        raise HTTPException(status_code=409, detail="Only locked payroll can be exported")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payroll"
    headers = [
        "Employee code",
        "Employee name",
        "Currency",
        "Base pay",
        "Overtime pay",
        "Allowances",
        "Bonus",
        "Deductions",
        "Unpaid leave",
        "Gross pay",
        "Net pay",
        "Payment status",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="123047")

    items = (
        db.query(PayrollItem, Employee)
        .join(Employee, PayrollItem.employee_id == Employee.id)
        .filter(PayrollItem.run_id == run.id)
        .order_by(Employee.employee_code)
        .all()
    )
    for item, employee in items:
        sheet.append(
            [
                employee.employee_code,
                employee.full_name,
                run.currency,
                float(item.base_pay),
                float(item.overtime_pay),
                float(item.allowances),
                float(item.bonus_pay),
                float(item.deductions),
                float(item.unpaid_leave_deduction),
                float(item.gross_pay),
                float(item.net_pay),
                item.payment_status,
            ]
        )
        for cell in list(sheet[sheet.max_row][:3]) + [sheet.cell(sheet.max_row, 12)]:
            cell.data_type = "s"
        for cell in sheet[sheet.max_row][3:11]:
            cell.number_format = f'"{run.currency}" #,##0.00'

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {
        "A": 18,
        "B": 30,
        "C": 12,
        "D": 16,
        "E": 16,
        "F": 16,
        "G": 16,
        "H": 16,
        "I": 16,
        "J": 16,
        "K": 16,
        "L": 18,
    }.items():
        sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    record_audit(
        db,
        current_user,
        "payroll.xlsx_exported",
        "payroll_run",
        run.id,
        company_id=run.company_id,
    )
    db.commit()
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="payroll-{run.id}.xlsx"'
        },
    )
